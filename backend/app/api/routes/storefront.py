import asyncio
import hashlib
import io
import json
import re
import secrets
import time
from datetime import datetime

import httpx
from fastapi import APIRouter, HTTPException, Query, Request, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

from app.core.config import settings
from app.db import user_store
from app.services.deepseek_client import deepseek_client
from app.services.opencart_api import opencart_api
from app.services.product_db import product_db
from app.services.catalog_core import catalog_core, is_doc_header_line
from app.services.dadata_api import dadata_api
from app.services.invoice_generator import render_invoice_pdf

# Публичный каталог для гео-витрин (moscow.stroiapp.ru и др.).
# Читается из ProductDB — того же источника, что и раздел «ТОВАРЫ НА САЙТЕ».
# Отдаёт ТОЛЬКО безопасные поля: без себестоимости, опта, ROI и служебных метрик.
router = APIRouter(prefix="/storefront", tags=["storefront"])
# Алиас v1-API: тот же пайплайн, короткий путь /api/v1/parser/upload.
v1_router = APIRouter(prefix="/parser", tags=["parser-v1"])

SHOP = "https://stroiapp.ru"
IMG = SHOP + "/image/"


def _img(image: str) -> str:
    if not image:
        return ""
    s = str(image)
    return s if s.startswith("http") else IMG + s


# Фото разделов лежат на stroysmatom.ru (источник импорта каталога) — на stroiapp.ru их нет.
MIKS_IMG = "https://stroysmatom.ru/upload/foto_razdel_price/"


def _cat_img(image: str) -> str:
    if not image:
        return ""
    s = str(image).rsplit("/", 1)[-1]          # basename
    s = re.sub(r"\(\d+\)(?=\.\w+$)", "", s)  # name(1).jpg -> name.jpg
    return MIKS_IMG + s


def _product_url(pid) -> str:
    return f"{SHOP}/index.php?route=product/product&product_id={pid}"


# ---------- Журнал moscow.stroiapp.ru (контент-фабрика) ----------

def _journal_card(a: dict) -> dict:
    return {
        "slug": a["slug"],
        "title": a["title"],
        "category": a["category"],
        "tone": a["tone"],
        "intro": a["intro"],
        "read": a["read_time"],
        "product_sku": a["product_sku"],
        "published_at": a["published_at"],
    }


@router.get("/journal")
def storefront_journal(limit: int = Query(20, ge=1, le=50)):
    """Опубликованные статьи Журнала — карточки для секции на витрине."""
    from app.db import journal_store
    items = journal_store.list_articles(status="published", limit=min(limit, 50))
    return {"articles": [_journal_card(a) for a in items], "total": len(items)}


@router.get("/journal/{slug}")
def storefront_journal_article(slug: str):
    """Полная статья Журнала по slug."""
    from app.db import journal_store
    a = journal_store.get_by_slug(slug, published_only=True)
    if not a:
        return {"article": None}
    return {"article": {
        **_journal_card(a),
        "content": a["content"],
        "product_name": a["product_name"],
    }}


# ---------- Счётчик просмотров Журнала (кто читал, из какой страны) ----------

_geo_cache: dict[str, tuple[str, str]] = {}  # ip -> (country, city)


async def _geo_ip(ip: str) -> tuple[str, str]:
    """Страна+город по IP через ip-api.com (free, ~45 req/min). Кэшируем в памяти."""
    if (not ip or ip in ("127.0.0.1", "::1")
            or ip.startswith(("192.168.", "10.", "169.254."))):
        return "", ""
    if ip in _geo_cache:
        return _geo_cache[ip]
    result = ("", "")
    try:
        async with httpx.AsyncClient(timeout=3.0) as client:
            r = await client.get(
                f"http://ip-api.com/json/{ip}?fields=status,country,city&lang=ru"
            )
            d = r.json()
        if d.get("status") == "success":
            result = (d.get("country") or "", d.get("city") or "")
    except Exception:
        pass
    if len(_geo_cache) > 5000:
        _geo_cache.clear()
    _geo_cache[ip] = result
    return result


async def _record_journal_view(slug: str, ip: str, ua: str, ref: str) -> None:
    """Фоновая запись просмотра: статья по slug → дедуп по IP → гео → journal_views."""
    from app.db import journal_store
    try:
        a = journal_store.get_by_slug(slug, published_only=True)
        if not a:
            return
        if journal_store.recent_view_exists(a["id"], ip):
            return
        country, city = await _geo_ip(ip)
        journal_store.record_view(a["id"], ip=ip, country=country, city=city,
                                  user_agent=ua, referrer=ref)
    except Exception as exc:
        print(f"[journal] view track failed for {slug}: {exc}")


@router.post("/journal/{slug}/view")
async def storefront_journal_view(slug: str, request: Request):
    """Счётчик просмотров статьи. Страница /journal?slug=... шлёт POST при открытии."""
    fwd = request.headers.get("x-forwarded-for", "")
    ip = fwd.split(",")[0].strip() or (request.client.host if request.client else "")
    ua = (request.headers.get("user-agent") or "")[:400]
    ref = (request.headers.get("referer") or "")[:400]
    asyncio.create_task(_record_journal_view(slug, ip, ua, ref))
    return {"status": "ok"}


def _public(p) -> dict:
    return {
        "id": p.product_id,
        "sku": p.sku,
        "name": p.name,
        "model": p.model,
        "price": p.retail_price,
        "image": _img(p.image),
        "quantity": p.quantity,
        "category_id": p.category_id,
        "category_ids": p.category_ids or [],
        "attributes": p.attributes or [],
        "description": p.description or "",
        "unit": catalog_core.facets(p.sku).get("unit", "шт"),
        "facets": catalog_core.facets(p.sku),
        "url": _product_url(p.product_id),
    }


@router.get("/products")
def storefront_products(limit: int = Query(48, ge=1, le=100),
                        category_id: int | None = None,
                        search: str = "", offset: int = Query(0, ge=0)):
    items = product_db.list_products()
    if category_id is not None:
        items = [p for p in items if p.category_id == category_id or category_id in (p.category_ids or [])]
    if search:
        q = search.lower()

        def _hit(p) -> bool:
            if q in (p.name or "").lower() or q in (p.sku or "").lower() or q in (p.model or "").lower():
                return True
            # бренд и прочие характеристики лежат в атрибутах [{name, text}]
            return any(q in str(a.get("text", "")).lower() for a in (p.attributes or []))

        items = [p for p in items if _hit(p)]
    out = [_public(p) for p in items if p.name and p.retail_price > 0]
    page = out[max(offset, 0):max(offset, 0) + limit]
    return {"status": "ok", "count": len(page), "total": len(out), "products": page}


@router.get("/product/{key}")
def storefront_product(key: str):
    """Карточка товара по product_id или sku — для страницы товара на витрине."""
    items = product_db.list_products()
    p = None
    if key.isdigit():
        pid = int(key)
        p = next((x for x in items if x.product_id == pid), None)
    if p is None:
        p = product_db.get_by_sku(key)
    if p is None:
        p = next((x for x in items if str(x.sku) == key or str(x.product_id) == key), None)
    if not p or not p.name:
        raise HTTPException(status_code=404, detail="Товар не найден")
    return {"status": "ok", "product": _public(p)}


_cat_cache = {"ts": 0.0, "data": []}


@router.get("/categories")
async def storefront_categories():
    # кэш 5 мин — категории меняются редко, а OpenCart дергать на каждый хит не надо
    if _cat_cache["data"] and time.time() - _cat_cache["ts"] < 300:
        return {"status": "ok", "categories": _cat_cache["data"]}
    data = await opencart_api.get_action("category/list")
    arr = (data or {}).get("categories") or []

    # счётчики товаров по каждой категории любого уровня (прямые привязки + верхняя)
    counts: dict[int, int] = {}
    for p in product_db.list_products():
        if not (p.name and p.retail_price > 0):
            continue
        seen = set(p.category_ids or [])
        if p.category_id:
            seen.add(p.category_id)
        for cid in seen:
            counts[cid] = counts.get(cid, 0) + 1

    def node(c) -> dict | None:
        cid = str(c.get("category_id", ""))
        if not cid.isdigit() or not c.get("name"):
            return None
        return {
            "category_id": int(cid),
            "name": c.get("name"),
            "image": _cat_img(c.get("image")),
            "count": counts.get(int(cid), 0),
        }

    by_parent: dict[int, list] = {}
    for c in arr:
        n = node(c)
        if n:
            by_parent.setdefault(int(c.get("parent_id") or 0), []).append(n)

    top = by_parent.get(0, [])
    for t in top:
        children = by_parent.get(t["category_id"], [])
        children.sort(key=lambda c: -c["count"])
        t["children"] = children
        # если у верхней категории нет прямых товаров — счётчик = сумма детей
        if not t["count"]:
            t["count"] = sum(ch["count"] for ch in children)
    top.sort(key=lambda c: -c["count"])
    if top:
        _cat_cache.update({"ts": time.time(), "data": top})
    return {"status": "ok", "categories": top}


# ---------- ИИ-калькулятор: данные СТРОГО из раздела «Товары на сайте» (ProductDB) ----------

class CalcRequest(BaseModel):
    room_type: str = ""
    work_type: str = ""
    base_type: str = ""
    area: float = 0
    thickness: float = 0
    comment: str = ""
    query: str = ""  # свободный текст из умной строки ИИ-калькулятора


_CALC_SYSTEM = (
    "Ты — ИИ-калькулятор маркетплейса StroiApp. Твоя задача — рассчитать материалы "
    "по формулам расхода и предложить товары.\n"
    "КРИТИЧЕСКОЕ ТРЕБОВАНИЕ: тебе передан динамический список из нашей админки "
    "«Товары на сайте». Ты имеешь право рекомендовать и выводить ТОЛЬКО те товары, "
    "которые есть в этом списке прямо сейчас (названия, артикулы, цены берёшь строго "
    "оттуда). Если товар удалён из этого раздела админки или его нет в наличии — ты ни "
    "в коем случае не должен его выдумывать и предлагать. Если позиция только что "
    "добавлена — сразу используй её в расчётах.\n"
    "Отвечай СТРОГО валидным JSON без markdown:\n"
    '{"summary": "краткий текстовый расчёт для клиента (2-4 предложения, по-русски)", '
    '"items": [{"sku": "артикул товара ИЗ СПИСКА", "qty": число, "unit": "шт|мешков|кг|л|м2", '
    '"note": "зачем нужен, 3-8 слов"}]}\n'
    "Правила: qty — целое число упаковок/единиц с округлением вверх; выбирай 2-6 самых "
    "подходящих товаров; sku копируй из списка посимвольно; если подходящих товаров нет — "
    'верни {"summary": "объяснение", "items": []}.'
)


@router.post("/ai-calc")
async def storefront_ai_calc(req: CalcRequest):
    """ИИ-расчёт материалов: контекст = актуальный список из ProductDB («Товары на сайте»).
    Принимает либо структурированные поля (area обязательна), либо свободный query."""
    q = (req.query or "").strip()
    if req.area <= 0 and not q:
        return {"status": "error", "detail": "area or query required"}

    items = [p for p in product_db.list_products() if p.name and p.retail_price > 0]
    if not items:
        return {"status": "error", "detail": "catalog empty"}

    # Компактный JSON-lines каталог для контекста DeepSeek
    catalog = "\n".join(
        json.dumps(
            {
                "sku": p.sku,
                "name": p.name,
                "price": p.retail_price,
                "stock": "в наличии" if (p.quantity or 0) > 0 else "под заказ",
            },
            ensure_ascii=False,
        )
        for p in items[:400]
    )

    user = (
        "Параметры задачи клиента:\n"
        + (f"- Запрос клиента своими словами: {q}\n" if q else "")
        + f"- Тип помещения: {req.room_type or 'не указан'}\n"
        f"- Тип работ: {req.work_type or 'не указан'}\n"
        f"- Тип основания: {req.base_type or 'не указан'}\n"
        f"- Площадь: {req.area} м2\n"
        f"- Толщина слоя: {req.thickness} мм\n"
        + (f"- Комментарий: {req.comment}\n" if req.comment else "")
        + ("\nКлиент описал задачу свободным текстом — сам извлеки из него тип работ, "
           "площадь, толщину и прочие параметры; если площадь не названа — предположи "
           "типичную и оговори предположение в summary.\n" if q else "")
        + "\nАктуальный список товаров из админки «Товары на сайте» "
        "(JSON-строки: sku, name, price в рублях, stock):\n"
        + catalog
    )

    raw = await deepseek_client.chat(user, max_tokens=1500, temperature=0.3, system=_CALC_SYSTEM)
    if not raw or raw.startswith("[DeepSeek error"):
        return {"status": "error", "detail": raw or "deepseek unavailable"}

    # Парсим JSON-ответ модели (терпимо к обёрткам ```json ... ```)
    text = raw.strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:json)?\s*", "", text)
        text = re.sub(r"\s*```$", "", text)
    data = None
    try:
        data = json.loads(text)
    except ValueError:
        m = re.search(r"\{.*\}", text, re.DOTALL)
        if m:
            try:
                data = json.loads(m.group(0))
            except ValueError:
                data = None

    if not isinstance(data, dict):
        # Модель ответила текстом — отдаём как есть, без карточек
        return {"status": "ok", "text": raw, "items": []}

    # Матчим sku из ответа на реальные товары — выдуманные позиции отсекаются здесь
    by_sku = {str(p.sku): p for p in items}
    out_items = []
    for it in data.get("items") or []:
        p = by_sku.get(str(it.get("sku", "")).strip())
        if not p:
            continue
        card = _public(p)
        try:
            card["qty"] = max(1, int(float(it.get("qty") or 1)))
        except (TypeError, ValueError):
            card["qty"] = 1
        card["unit"] = str(it.get("unit") or "шт")
        card["note"] = str(it.get("note") or "")
        out_items.append(card)

    return {"status": "ok", "text": str(data.get("summary") or ""), "items": out_items}


# ---------- ИИ-расчёт по файлу спецификации (Excel / PDF / Word) ----------

_SPEC_SYSTEM = (
    "Ты — закупщик-ИИ маркетплейса StroiApp. Клиент прислал файл со спецификацией "
    "(ведомостью материалов). Твоя задача — распознать позиции и подобрать для каждой "
    "товар из нашего КАТАЛОГА (раздел админки «Товары на сайте»).\n"
    "КРИТИЧЕСКОЕ ТРЕБОВАНИЕ: предлагай ТОЛЬКО товары из переданного списка каталога. "
    "Аналог допустим только если он реально заменяет позицию по назначению "
    "(например, другая грунтовка вместо указанной) — тогда помечай в note «аналог». "
    "Если в каталоге нет ни товара, ни осмысленной замены (например, крепёж, лента, "
    "метизы, которых нет в списке) — НЕ подбирай случайный товар: добавь позицию "
    "в массив not_found. Никогда не выдумывай товары и артикулы.\n"
    "Отвечай СТРОГО валидным JSON без markdown:\n"
    '{"summary": "1-2 предложения: что распознано из спецификации, по-русски", '
    '"items": [{"sku": "артикул ИЗ КАТАЛОГА", "qty": число, "unit": "шт|мешков|кг|л|м2", '
    '"note": "позиция из спецификации, 2-6 слов"}], '
    '"not_found": ["точная формулировка позиции из спецификации"]}\n'
    "Правила: qty — количество упаковок с округлением ВВЕРХ (пересчитай вес/объём из "
    "спецификации в упаковки по фасовке товара); позиций может быть сколько угодно "
    "(обычно 3-15); sku копируй из каталога посимвольно; каждая позиция спецификации "
    "должна попасть либо в items, либо в not_found; если ничего не подобралось — "
    'верни {"summary": "объяснение", "items": [], "not_found": [...]}.'
)


def _extract_spec_text(filename: str, blob: bytes) -> str:
    """Текст спецификации из файла. Возвращает пустую строку, если не распознали."""
    fn = filename.lower()

    if fn.endswith((".xlsx", ".xlsm")):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
        lines = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                cells = [str(c).strip() for c in row if c not in (None, "")]
                if cells:
                    lines.append(" | ".join(cells))
        return "\n".join(lines)

    if fn.endswith(".docx"):
        import docx
        d = docx.Document(io.BytesIO(blob))
        lines = [p.text.strip() for p in d.paragraphs if p.text.strip()]
        for tbl in d.tables:
            for row in tbl.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    lines.append(" | ".join(cells))
        return "\n".join(lines)

    if fn.endswith(".pdf"):
        from pypdf import PdfReader
        r = PdfReader(io.BytesIO(blob))
        lines = []
        for page in r.pages[:15]:
            t = (page.extract_text() or "").strip()
            if t:
                lines.append(t)
        return "\n".join(lines)

    return ""


async def _match_spec_with_ai(spec_text: str, source: str) -> dict:
    """Общая связка OCR/парсинга с DeepSeek: текст спецификации → смета из живого
    каталога. Ядро catalog_core отбирает релевантное подмножество кандидатов —
    модель матчит по ~120 точным позициям, а не по слепым первым 400 записям."""
    products = [p for p in product_db.list_products() if p.name and p.retail_price > 0]
    if not products:
        return {"status": "error", "detail": "catalog empty"}

    # Шапка/подвал счёта или накладной (поставщик, ИНН, НДС…) — не позиции
    body_lines = [ln for ln in spec_text.splitlines()
                  if not is_doc_header_line(ln)]
    clean_text = "\n".join(body_lines) or spec_text

    # Кандидаты под эту спецификацию (нормализация, синонимы, фасовка, обучение)
    candidates = catalog_core.relevant_for_spec(clean_text, limit=120)
    candidates = [p for p in candidates if p.retail_price > 0] or products[:120]

    def _card(p) -> dict:
        d = {
            "sku": p.sku,
            "name": p.name,
            "price": p.retail_price,
            "stock": "в наличии" if (p.quantity or 0) > 0 else "под заказ",
        }
        f = catalog_core.facets(p.sku)
        if f.get("pack"):
            d["pack"] = f["pack"]
        if f.get("frost"):
            d["frost"] = f["frost"]
        if f.get("scope"):
            d["scope"] = f["scope"]
        if f.get("color"):
            d["color"] = f["color"]
        return d

    catalog = "\n".join(
        json.dumps(_card(p), ensure_ascii=False) for p in candidates
    )

    # few-shot: выученные связи «фраза спеки → sku» подмешиваем в промпт
    try:
        from app.db import learn_store as _ls
        shots = _ls.top_aliases(15)
    except Exception:
        shots = []
    learned_hint = ""
    if shots:
        learned_hint = (
            "\n\nВыученные ранее соответствия (позиция спеки → артикул, "
            "применяй их в первую очередь):\n"
            + "\n".join(f'"{s["phrase"]}" -> {s["sku"]}' for s in shots)
        )

    user = (
        f"Спецификация клиента (источник: {source}; строки могут содержать "
        "название, единицу, количество, вес; возможны ошибки OCR — исправляй их по смыслу):\n"
        "---\n" + clean_text + "\n---\n"
        "\nАктуальный список товаров из админки «Товары на сайте» "
        "(JSON-строки: sku, name, price в рублях, stock):\n" + catalog
        + learned_hint
    )

    raw = await deepseek_client.chat(user, max_tokens=2500, temperature=0.2, system=_SPEC_SYSTEM)
    if not raw or raw.startswith("[DeepSeek error"):
        return {"status": "error", "detail": raw or "deepseek unavailable"}

    text = raw.strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:json)?\s*", "", text)
        text = re.sub(r"\s*```$", "", text)
    data = None
    try:
        data = json.loads(text)
    except ValueError:
        m = re.search(r"\{.*\}", text, re.DOTALL)
        if m:
            try:
                data = json.loads(m.group(0))
            except ValueError:
                data = None

    if not isinstance(data, dict):
        return {"status": "ok", "file": source, "text": raw, "items": []}

    by_sku = {str(p.sku): p for p in products}
    out_items = []
    for it in data.get("items") or []:
        p = by_sku.get(str(it.get("sku", "")).strip())
        if not p:
            continue
        card = _public(p)
        try:
            card["qty"] = max(1, int(float(it.get("qty") or 1)))
        except (TypeError, ValueError):
            card["qty"] = 1
        card["unit"] = str(it.get("unit") or "шт")
        card["note"] = str(it.get("note") or "")
        out_items.append(card)

        # авто-дообучение: связь «позиция спеки → sku» усиливает ядро
        if card["note"]:
            try:
                catalog_core.learn(card["note"], p.sku, True, source="auto")
            except Exception:
                pass

    not_found = [str(x).strip() for x in (data.get("not_found") or []) if str(x).strip()]
    # Ядро запоминает спрос: позиции, которых нет в каталоге
    for t in not_found:
        try:
            catalog_core.observe_not_found(t, source=source)
        except Exception:
            pass
    # Связанные товары: позиции одного документа — связи в product_link
    try:
        catalog_core.observe_bundle([c["sku"] for c in out_items],
                                    source=source)
    except Exception:
        pass
    return {"status": "ok", "file": source,
            "text": str(data.get("summary") or ""),
            "items": out_items, "not_found": not_found}


@router.post("/ai-calc-file")
async def storefront_ai_calc_file(file: UploadFile):
    """Спецификация из Excel/PDF/Word → смета из живого каталога ProductDB."""
    name = file.filename or "spec"
    blob = await file.read()
    if len(blob) > 8 * 1024 * 1024:
        return {"status": "error", "detail": "Файл слишком большой (максимум 8 МБ)"}
    if len(blob) < 100:
        return {"status": "error", "detail": "Файл выглядит пустым"}

    # Фото и скан-PDF — сразу в цепочку компьютерного зрения.
    if name.lower().endswith((".jpg", ".jpeg", ".png", ".webp", ".bmp")):
        return await _parser_pipeline(name, blob)

    try:
        spec_text = _extract_spec_text(name, blob)
    except Exception as exc:
        print(f"[ai-calc-file] extract failed for {name}: {exc}")
        spec_text = ""

    # Скан без текстового слоя (PDF как картинка) — шлём в CV-цепочку.
    if name.lower().endswith(".pdf") and len(spec_text.strip()) < 10:
        return await _parser_pipeline(name, blob)

    if len(spec_text.strip()) < 10:
        return {"status": "error",
                "detail": "Не удалось распознать текст в файле (форматы: .xlsx, .docx, .pdf)"}

    if len(spec_text) > 12000:
        spec_text = spec_text[:12000] + "\n…(обрезано)"

    return await _match_spec_with_ai(spec_text, name)

# ---------- Компьютерное зрение: фото/сканы спецификаций (EasyOCR → DeepSeek) ----------
# Стек под Timeweb 2 vCPU / 4 ГБ, CPU-only:
#  - easyocr (torch CPU): лучше всего читает русский рукописный текст и строительные
#    спецификации; Reader создаётся лениво ОДИН раз (модели ~100 МБ) и держится в памяти.
#  - tesseract (multi-pass с бинаризацией Оцу) остаётся фолбэком, если easyocr
#    недоступен/не распознал ничего.
#  - OCR блокирующий: гоняется через asyncio.to_thread, чтобы не вешать event loop.

_EASY_READER = None
_EASY_BROKEN = False


def _get_easy_reader():
    """Ленивая инициализация EasyOCR: ru+en, строго CPU, 2 потока (2 vCPU)."""
    global _EASY_READER, _EASY_BROKEN
    if _EASY_READER is not None or _EASY_BROKEN:
        return _EASY_READER
    try:
        import torch
        torch.set_num_threads(2)  # у Timeweb 2 vCPU — не даём torch съесть всё
        import easyocr
        _EASY_READER = easyocr.Reader(["ru", "en"], gpu=False, verbose=False)
        print("[parser/ocr] EasyOCR инициализирован (CPU, ru+en)")
    except Exception as exc:
        _EASY_BROKEN = True
        print(f"[parser/ocr] EasyOCR недоступен ({exc}) — будет Tesseract")
    return _EASY_READER


def _easy_read_sync(blob: bytes) -> str:
    """EasyOCR по байтам картинки → строки текста. Пустая строка = не распознал."""
    reader = _get_easy_reader()
    if reader is None:
        return ""
    try:
        import numpy as np
        from PIL import Image
        img = Image.open(io.BytesIO(blob)).convert("RGB")
        # апскейл мелких фото — русские модели easyocr лучше читают от ~1600px
        if max(img.size) < 1600:
            f = 1600 / max(img.size)
            img = img.resize((int(img.width * f), int(img.height * f)),
                             Image.Resampling.LANCZOS)
        texts = reader.readtext(np.asarray(img), detail=0, paragraph=False,
                                min_size=8, low_text=0.3, text_threshold=0.5)
        return "\n".join(t.strip() for t in texts if t and t.strip())
    except Exception as exc:
        print(f"[parser/ocr] easyocr failed: {exc}")
        return ""


def _otsu_threshold(gray) -> int:
    """Порог бинаризации Оцу по гистограмме (без numpy/cv2)."""
    hist = gray.histogram()
    total = sum(hist)
    sum_all = sum(i * hist[i] for i in range(256))
    sum_b = w_b = best_t = best_var = 0
    for t in range(256):
        w_b += hist[t]
        if not w_b:
            continue
        w_f = total - w_b
        if not w_f:
            break
        sum_b += t * hist[t]
        m_b = sum_b / w_b
        m_f = (sum_all - sum_b) / w_f
        var = w_b * w_f * (m_b - m_f) ** 2
        if var > best_var:
            best_var, best_t = var, t
    return best_t


def _ocr_score(img, cfg: str) -> tuple[int, str]:
    """Один прогон tesseract: возвращает (score, text). Score = сумма conf слов.
    Слова собираются обратно в строки по block/par/line, чтобы модель видела
    структуру списка (строка спецификации = одна позиция)."""
    import pytesseract
    from pytesseract import Output

    data = pytesseract.image_to_data(img, lang="rus+eng", config=cfg,
                                     output_type=Output.DICT)
    lines: dict[tuple, list[str]] = {}
    conf_sum = 0.0
    for i in range(len(data["text"])):
        t = (data["text"][i] or "").strip()
        try:
            c = float(data["conf"][i])
        except (TypeError, ValueError):
            c = -1
        if not t or c < 0:
            continue
        key = (data["block_num"][i], data["par_num"][i], data["line_num"][i])
        lines.setdefault(key, []).append(t)
        conf_sum += c
    text = "\n".join(" ".join(ws) for ws in lines.values())
    return int(conf_sum), text
def _tesseract_read_sync(blob: bytes) -> str:
    """Фолбэк-Tesseract: 3 прогона (серый psm6, Оцу psm6, серый psm11), выбор по conf."""
    try:
        import pytesseract  # noqa: F401
        from PIL import Image, ImageFilter, ImageOps
    except ImportError:
        print("[parser/ocr] pytesseract/Pillow не установлены")
        return ""
    try:
        img = Image.open(io.BytesIO(blob)).convert("L")
        # апскейл до 2000px — телефонные фото резко лучше читаются
        if max(img.size) < 2000:
            f = 2000 / max(img.size)
            img = img.resize((int(img.width * f), int(img.height * f)),
                             Image.Resampling.LANCZOS)
        gray = ImageOps.autocontrast(img, cutoff=1).filter(ImageFilter.SHARPEN)
        bin_img = gray.point(
            lambda p, t=_otsu_threshold(gray): (255 if p > t else 0))

        candidates = [
            _ocr_score(gray, "--oem 1 --psm 6"),
            _ocr_score(bin_img, "--oem 1 --psm 6"),
            _ocr_score(gray, "--oem 1 --psm 11"),
        ]
        best = max(candidates, key=lambda c: c[0])
        print(f"[parser/ocr] tesseract scores={[c[0] for c in candidates]}")
        return best[1].strip()
    except Exception as exc:
        print(f"[parser/ocr] tesseract failed: {exc}")
        return ""


async def _ocr_image(blob: bytes) -> str:
    """Сначала EasyOCR (рукопись/русский), при провале — multi-pass Tesseract."""
    text = await asyncio.to_thread(_easy_read_sync, blob)
    if len(text.strip()) >= 10:
        return text.strip()
    print("[parser/ocr] easyocr пустой/короткий — fallback на tesseract")
    return await asyncio.to_thread(_tesseract_read_sync, blob)


async def _ocr_pdf_scan(blob: bytes) -> str:
    """Скан-PDF → растеризация страниц (PyMuPDF 200dpi) → OCR каждой страницы."""
    try:
        import fitz  # PyMuPDF
    except ImportError:
        print("[parser/ocr] PyMuPDF не установлен — сканы PDF не умею")
        return ""
    try:
        doc = fitz.open(stream=blob, filetype="pdf")
    except Exception as exc:
        print(f"[parser/ocr] pdf open failed: {exc}")
        return ""
    parts = []
    try:
        for page in doc[:8]:  # не больше 8 страниц за раз
            pix = page.get_pixmap(dpi=200)
            parts.append(await _ocr_image(pix.tobytes("png")))
    except Exception as exc:
        print(f"[parser/ocr] pdf raster failed: {exc}")
    finally:
        doc.close()
    return "\n".join(p for p in parts if p).strip()


async def _parser_pipeline(name: str, blob: bytes) -> dict:
    """OCR фото/скана → DeepSeek → смета. Вызывается и роутом, и ai-calc-file."""
    fn = name.lower()
    if fn.endswith((".jpg", ".jpeg", ".png", ".webp", ".bmp")):
        text = await _ocr_image(blob)
    elif fn.endswith(".pdf"):
        text = await _ocr_pdf_scan(blob)
    else:
        return {"status": "error",
                "detail": "Компьютерное зрение принимает фото (.jpg, .png) и PDF-сканы"}

    if len(text.strip()) < 10:
        return {"status": "error",
                "detail": "Текст на изображении не распознан — попробуйте фото чётче "
                          "и хорошо освещённым"}
    if len(text) > 12000:
        text = text[:12000] + "\n…(обрезано)"

    print(f"[parser/ocr] {name}: распознано {len(text)} символов")
    return await _match_spec_with_ai(text, name)


@router.post("/parser/upload")
async def storefront_parser_upload(file: UploadFile):
    """Фото рукописного списка / скан спецификации (.jpg/.png/.pdf) → OCR → смета.
    Отвечает тем же JSON, что и /ai-calc-file — фронт рендерит единую смету."""
    name = file.filename or "scan"
    blob = await file.read()
    if len(blob) > 8 * 1024 * 1024:
        return {"status": "error", "detail": "Файл слишком большой (максимум 8 МБ)"}
    if name.lower().endswith((".xlsx", ".xlsm", ".docx")):
        try:
            spec_text = _extract_spec_text(name, blob)
        except Exception:
            spec_text = ""
        if len(spec_text.strip()) < 10:
            return {"status": "error",
                    "detail": "Не удалось распознать текст в файле (.xlsx/.docx)"}
        return await _match_spec_with_ai(spec_text[:12000], name)
    return await _parser_pipeline(name, blob)


# ---------- Готовые комплекты для умного калькулятора ----------
# Курируемые наборы: собираются из актуальных товаров ProductDB по ключевым словам.
_KIT_DEFS = [
    {
        "id": "plaster",
        "title": "Штукатурка + грунтовка",
        "note": "комплект для сухих комнат ~20 м²",
        "items": [
            {"keys": ["штукатурк", "ротбанд"], "qty": 5},
            {"keys": ["грунт"], "qty": 1},
        ],
    },
    {
        "id": "tile",
        "title": "Плиточный клей + затирка",
        "note": "комплект для санузлов ~8 м²",
        "items": [
            {"keys": ["плиточный клей", "клей для плитк", "клей плиточн"], "qty": 4},
            {"keys": ["затирк"], "qty": 1},
        ],
    },
    {
        "id": "partitions",
        "title": "Монтаж перегородок из ГКЛ",
        "note": "перегородка ~10 м²: гипсокартон, профили, саморезы",
        "items": [
            {"keys": ["гипсокартон", "гкл"], "qty": 6},
            {"keys": ["профиль"], "qty": 10},
            {"keys": ["саморез"], "qty": 1},
        ],
    },
    {
        "id": "floor",
        "title": "Устройство стяжки пола",
        "note": "комплект для пола ~15 м²",
        "items": [
            {"keys": ["стяжк", "пескобетон"], "qty": 8},
            {"keys": ["маяк"], "qty": 5},
            {"keys": ["демпферн", "кромочн", "лента"], "qty": 1},
        ],
    },
    {
        "id": "waterproof",
        "title": "Гидроизоляция санузла под ключ",
        "note": "обмазочная защита ~6 м²",
        "items": [
            {"keys": ["гидроизоляц", "мастик"], "qty": 1},
            {"keys": ["лента"], "qty": 1},
            {"keys": ["валик"], "qty": 1},
        ],
    },
    {
        "id": "putty",
        "title": "Финишная шпаклевка стен",
        "note": "подготовка под покраску ~40 м²",
        "items": [
            {"keys": ["шпаклевк", "шпатлевк"], "qty": 3},
            {"keys": ["стеклохолст", "серпянк", "сетк"], "qty": 1},
            {"keys": ["шпател"], "qty": 1},
        ],
    },
    {
        "id": "insulation",
        "title": "Утепление лоджии / балкона",
        "note": "~12 м² экструзионного утеплителя",
        "items": [
            {"keys": ["пеноплэкс", "пеноплекс", "экструзион", "пенополистирол"], "qty": 6},
            {"keys": ["клей-пена", "пена", "клей"], "qty": 2},
            {"keys": ["дюбель", "грибок", "зонтик"], "qty": 1},
        ],
    },
    {
        "id": "paint",
        "title": "Малярные работы (покраска)",
        "note": "стены ~50 м² в два слоя",
        "items": [
            {"keys": ["краск", "tikkurila", "тиккурил"], "qty": 2},
            {"keys": ["грунт"], "qty": 1},
            {"keys": ["валик"], "qty": 2},
            {"keys": ["скотч", "малярн"], "qty": 2},
        ],
    },
    {
        "id": "electric",
        "title": "Электрика: разводка на комнату",
        "note": "розетки + выключатели ~20 м²",
        "items": [
            {"keys": ["кабель", "ввг"], "qty": 3},
            {"keys": ["подрозетник"], "qty": 10},
            {"keys": ["гофр"], "qty": 2},
        ],
    },
    {
        "id": "warmfloor",
        "title": "Тёплый пол в стяжку",
        "note": "нагревательный мат ~4 м²",
        "items": [
            {"keys": ["нагревательн", "тёплый пол", "теплый пол"], "qty": 1},
            {"keys": ["терморегулятор", "термостат"], "qty": 1},
            {"keys": ["гофр"], "qty": 1},
        ],
    },
    {
        "id": "laminate",
        "title": "Укладка ламината под ключ",
        "note": "комната ~18 м²",
        "items": [
            {"keys": ["ламинат"], "qty": 8},
            {"keys": ["подложк"], "qty": 2},
            {"keys": ["скотч"], "qty": 1},
        ],
    },
    {
        "id": "blocks",
        "title": "Армирование и кладка блоков",
        "note": "стена ~15 м² из газоблоков",
        "items": [
            {"keys": ["газоблок", "газобетон", "газосиликат", "блок"], "qty": 30},
            {"keys": ["клей"], "qty": 3},
            {"keys": ["сетк"], "qty": 1},
        ],
    },
]


@router.get("/core/status")
def storefront_core_status():
    """Health ядра каталога: размер индекса, токены, фасеты, обучение."""
    stats = catalog_core.stats()
    probe = catalog_core.search("ротбанд штукатурка 30кг", limit=3)
    from app.db import learn_store
    return {"status": "ok", "index": stats,
            "learning": learn_store.learn_stats(),
            "domains": learn_store.domain_stats(),
            "probe": [{"sku": p.sku, "name": p.name[:60]} for p in probe]}


@router.get("/core/knowledge")
def storefront_core_knowledge():
    """Карта знаний ядра: все разделы сервера, статические + выученные термины."""
    return {"status": "ok", "domains": catalog_core.knowledge()}


@router.get("/core/classify")
def storefront_core_classify(text: str = ""):
    """К какому разделу относится фраза — ядро маршрутизирует запрос."""
    return {"status": "ok", "text": text,
            "domain": catalog_core.domain_of(text),
            "ranked": catalog_core.classify(text)}


class _DomainLearn(BaseModel):
    domain: str = ""    # раздел: tenders/logistics/ads/geo/…
    term: str = ""      # термин или фраза раздела
    target: str = ""    # привязка: sku/город/действие (опционально)
    kind: str = "term"  # term|alias|fact|rule


@router.post("/core/learn")
async def core_domain_learn(payload: _DomainLearn):
    """Обучение ядра по разделам: термин привязывается к домену навсегда."""
    if not payload.domain.strip() or not payload.term.strip():
        return {"status": "error", "detail": "domain и term обязательны"}
    catalog_core.learn_domain(payload.domain, payload.term,
                              payload.target, payload.kind)
    return {"status": "ok", "learned": True, "domain": payload.domain}


@router.post("/core/retrain")
def storefront_core_retrain():
    """Принудительное переобучение: полная пересборка индекса ядра
    (каталог + learned_alias + rejected + domain_knowledge)."""
    stats = catalog_core.retrain()
    return {"status": "ok", "retrained": True, "index": stats}


@router.get("/core/product/{sku}")
def storefront_core_product(sku: str):
    """Досье товара в ядре: карточка, фасеты (бренд/категория/фасовка/цвет/
    морозостойкость), выученные фразы, отклонённые связи, связанные товары."""
    profile = catalog_core.product_profile(sku)
    if profile.get("status") == "error":
        return profile
    return {"status": "ok", "product": profile}


class _Feedback(BaseModel):
    phrase: str = ""   # формулировка позиции из спецификации (или название)
    sku: str = ""      # артикул товара, к которому относится вердикт
    accepted: bool = True
    domain: str = "product"  # 'product' → learned_alias/rejected; иначе → domain_knowledge


@router.post("/parser/feedback")
async def parser_feedback(payload: _Feedback):
    """Обучение ядра: пользователь принял/удалил позицию сметы.
    domain != 'product' пишется в доменную память раздела."""
    if not payload.phrase.strip():
        return {"status": "error", "detail": "phrase обязательна"}
    if payload.domain != "product":
        catalog_core.learn_domain(payload.domain, payload.phrase,
                                  payload.sku, "term", "feedback")
        return {"status": "ok", "learned": True, "domain": payload.domain}
    if not payload.sku.strip():
        return {"status": "error", "detail": "sku обязателен для product"}
    catalog_core.learn(payload.phrase, payload.sku, payload.accepted)
    return {"status": "ok", "learned": payload.accepted}


@router.get("/kits")
def storefront_kits():
    """Быстрые готовые комплекты — итоговая цена из актуальных товаров админки.
    Если хоть одна позиция комплекта не нашлась в каталоге — комплект не показываем."""
    products = [p for p in product_db.list_products() if p.name and p.retail_price > 0]

    def pick(keys: list[str]):
        hits = [p for p in products if any(k in p.name.lower() for k in keys)]
        if not hits:
            return None
        # в наличии — вперёд, внутри — по возрастанию цены (выгодный вариант)
        hits.sort(key=lambda p: (0 if (p.quantity or 0) > 0 else 1, p.retail_price))
        return hits[0]

    kits = []
    for kd in _KIT_DEFS:
        items = []
        for spec in kd["items"]:
            p = pick(spec["keys"])
            if not p:
                break
            card = _public(p)
            card["qty"] = spec["qty"]
            items.append(card)
        if len(items) != len(kd["items"]):
            continue
        total = round(sum(i["price"] * i["qty"] for i in items))
        kits.append({"id": kd["id"], "title": kd["title"], "note": kd["note"], "total": total, "items": items})

    # Комплекты из раздела «Акции» админки — promotion_kits (site_owner='moscow')
    try:
        for kd in user_store.list_kits("moscow", active_only=True):
            items = []
            for spec in kd["items"]:
                keys = spec.get("keys") or ([spec["name"]] if spec.get("name") else [])
                p = pick([str(k).lower() for k in keys]) if keys else None
                if not p:
                    break
                card = _public(p)
                card["qty"] = int(spec.get("qty") or 1)
                items.append(card)
            if items and len(items) == len(kd["items"]):
                kits.append({"id": f"db-{kd['id']}", "title": kd["title"], "note": kd.get("note", ""),
                             "total": round(sum(i["price"] * i["qty"] for i in items)),
                             "items": items, "site_owner": kd.get("site_owner", "moscow")})
    except Exception as exc:
        print("[storefront] db kits error:", exc)
    return {"status": "ok", "kits": kits}


# ---------- Голосовой ИИ-помощник: база знаний = раздел «Товары на сайте» (ProductDB) ----------

class ChatMsg(BaseModel):
    role: str = ""
    text: str = ""


class CartItem(BaseModel):
    sku: str = ""
    name: str = ""
    qty: int = 1
    price: float = 0


class OrderForm(BaseModel):
    name: str = ""
    phone: str = ""
    address: str = ""
    payment: str = ""


class ChatReq(BaseModel):
    message: str = ""
    history: list[ChatMsg] = []
    cart: list[CartItem] = []
    form: OrderForm | None = None


_VOICE_SYSTEM = (
    "Ты — голосовой ассистент маркетплейса StroiApp. Твоя база знаний — это "
    "динамический каталог из раздела админки «Товары на сайте». Отвечай голосом "
    "очень коротко и профессионально (1-3 предложения, без markdown и списков — "
    "текст будет озвучен).\n"
    "Помни контекст диалога: не переспрашивай то, что клиент уже сказал. "
    "Если клиент называет количество, город доставки и просит оформить — "
    "подтверди заказ: повтори товар, количество, город и итоговую сумму, "
    "и скажи, что открываешь корзину для оформления.\n"
    "Ты должен идеально ориентироваться в структуре наших категорий "
    "(Строительные материалы, Клей, Сантехника, Крепёж, Пиломатериалы, Инструмент, "
    "Лакокрасочные материалы и др.). Предлагай клиенту ТОЛЬКО те бренды, названия "
    "и цены, которые заведены в админку на текущую секунду. Если клиент ищет товар, "
    "которого нет в списке (удалён или никогда не был), прямо скажи, что позиции "
    "нет в наличии, и предложи точную альтернативу из доступных товаров.\n"
    "Отвечай СТРОГО валидным JSON без markdown:\n"
    '{"text": "короткий голосовой ответ клиенту", '
    '"items": [{"sku": "артикул товара ИЗ СПИСКА", "qty": 1}], '
    '"action": "none", "form": {"name": "", "phone": "", "address": "", "payment": ""}}\n'
    "items — до 3 самых релевантных товаров из списка, которые ты упомянул; "
    "если конкретных товаров не предлагаешь — items: []. sku копируй посимвольно.\n"
    "Ты видишь текущую корзину клиента (список в конце запроса) — отвечай на "
    "вопросы о её содержимом: называй позиции, количество и общую сумму.\n"
    "ВАЖНО: если клиент называет конкретный товар и количество — ВСЕГДА возвращай "
    "эти товары в items с точным qty, при любом action.\n"
    'action: "add" — клиент назвал конкретный товар и количество и просит '
    "добавить/положить в корзину («нужно 900 мешков», «добавь», «возьму») — "
    "верни эти товары в items с точным qty и action \"add\".\n"
    'action: "set" — клиент просит ИЗМЕНИТЬ количество позиции в корзине '
    "(«поставь 50», «сделай 10 мешков») ИЛИ называет количество товара, который "
    "уже есть в корзине — items с sku и итоговым qty.\n"
    'action: "remove" — клиент просит УДАЛИТЬ позицию из корзины — items с sku.\n'
    'action: "clear" — клиент просит очистить корзину.\n'
    'action: "reset" — клиент просит очистить/забыть историю разговора, '
    "начать диалог заново («очисти историю», «начнём сначала»).\n"
    'action: "cart" — клиент просит оформить заказ («оформляй», «оформить», '
    "«подтверждаю», «давай оформим», «всё, беру»), перейти в корзину или "
    "подтверждает покупку; иначе \"none\".\n"
    "Перед action \"cart\" проверь уже продиктованные данные заказа (список в "
    "конце запроса): если имя и телефон уже есть — НЕ переспрашивай их. "
    "Если имени или телефона нет — спроси недостающее голосом и верни action \"none\".\n"
    "form — данные заказа, которые клиент продиктовал голосом за весь диалог: "
    "имя, телефон (цифрами), адрес доставки, способ оплаты (наличные/карта/безнал). "
    "Неизвестные поля оставляй пустыми строками."
)


@router.post("/ai-chat")
async def storefront_ai_chat(req: ChatReq):
    """Голосовой помощник: контекст = актуальный срез ProductDB («Товары на сайте»)."""
    msg = (req.message or "").strip()
    if not msg:
        return {"status": "error", "detail": "message required"}

    items = [p for p in product_db.list_products() if p.name and p.retail_price > 0]
    if not items:
        return {"status": "error", "detail": "catalog empty"}

    catalog = "\n".join(
        json.dumps(
            {
                "sku": p.sku,
                "name": p.name,
                "price": p.retail_price,
                "stock": "в наличии" if (p.quantity or 0) > 0 else "под заказ",
            },
            ensure_ascii=False,
        )
        for p in items[:400]
    )

    history = ""
    for h in req.history[-8:]:
        who = "Клиент" if h.role == "user" else "Ассистент"
        t = (h.text or "").strip()
        if t:
            history += f"{who}: {t}\n"

    cart_txt = ""
    if req.cart:
        cart_txt = "Корзина клиента сейчас:\n" + "\n".join(
            f"- {c.name} (sku {c.sku}): {c.qty} шт × {c.price} ₽"
            for c in req.cart
        ) + "\n\n"

    form_txt = ""
    if req.form:
        parts = [
            f"имя: {req.form.name}" if req.form.name else "",
            f"телефон: {req.form.phone}" if req.form.phone else "",
            f"адрес: {req.form.address}" if req.form.address else "",
            f"оплата: {req.form.payment}" if req.form.payment else "",
        ]
        parts = [p for p in parts if p]
        if parts:
            form_txt = "Уже продиктованные данные заказа: " + ", ".join(parts) + "\n\n"

    user = (
        ("История диалога:\n" + history + "\n" if history else "")
        + "Вопрос клиента (распознан с голоса): " + msg + "\n\n"
        + cart_txt
        + form_txt
        + "Актуальный список товаров из админки «Товары на сайте» "
        + "(JSON-строки: sku, name, price в рублях, stock):\n" + catalog
    )

    raw = await deepseek_client.chat(user, max_tokens=700, temperature=0.4, system=_VOICE_SYSTEM)
    if not raw or raw.startswith("[DeepSeek error"):
        return {"status": "error", "detail": raw or "deepseek unavailable"}

    text = raw.strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:json)?\s*", "", text)
        text = re.sub(r"\s*```$", "", text)
    data = None
    try:
        data = json.loads(text)
    except ValueError:
        m = re.search(r"\{.*\}", text, re.DOTALL)
        if m:
            try:
                data = json.loads(m.group(0))
            except ValueError:
                data = None

    if not isinstance(data, dict):
        return {"status": "ok", "text": raw, "items": [], "action": "none", "form": {}}

    by_sku = {str(p.sku): p for p in items}
    out_items = []
    for it in data.get("items") or []:
        p = by_sku.get(str(it.get("sku", "")).strip())
        if not p:
            continue
        card = _public(p)
        try:
            card["qty"] = max(1, int(float(it.get("qty") or 1)))
        except (TypeError, ValueError):
            card["qty"] = 1
        out_items.append(card)

    action = str(data.get("action") or "").strip().lower()
    if action not in ("add", "set", "remove", "clear", "cart", "reset"):
        action = "none"
    raw_form = data.get("form") or {}
    form = {
        k: str(raw_form.get(k) or "").strip()
        for k in ("name", "phone", "address", "payment")
    } if isinstance(raw_form, dict) else {}

    return {"status": "ok", "text": str(data.get("text") or ""), "items": out_items, "action": action, "form": form}


# ---------- Реквизиты по ИНН + счёт на оплату (B2B) ----------

_INN_CACHE: dict[str, dict] = {}  # inn -> {"ts": float, "name": str}


@router.get("/geolocate")
async def storefront_geolocate(lat: float = 0, lon: float = 0):
    """Координаты → текстовый адрес через DaData geolocate/address.
    Фолбэк — Nominatim (OSM): бесплатный reverse-geocode без ключа.
    Публичный — для калькулятора доставки на витрине."""
    if not lat or not lon:
        return {"status": "error", "detail": "lat/lon обязательны"}
    address = await dadata_api.geolocate_address(lat, lon)
    if not address:
        address = await _nominatim_reverse(lat, lon)
    return {"status": "ok", "address": address}


async def _nominatim_reverse(lat: float, lon: float) -> str:
    """Бесплатный reverse-geocode OSM: собирает «город, улица, дом» по-русски."""
    try:
        async with httpx.AsyncClient(timeout=10, follow_redirects=True,
                                     headers={"User-Agent": "StroiApp/1.0 (moscow.stroiapp.ru)"}) as c:
            r = await c.get("https://nominatim.openstreetmap.org/reverse",
                            params={"format": "jsonv2", "lat": lat, "lon": lon,
                                    "zoom": 18, "accept-language": "ru"})
            d = r.json() or {}
        a = d.get("address") or {}
        road = a.get("road") or a.get("pedestrian") or a.get("residential") or ""
        house = a.get("house_number") or ""
        town = a.get("city") or a.get("town") or a.get("village") or a.get("suburb") or ""
        region = a.get("state") or ""
        parts = [p for p in [town, region if region != town else "", road,
                             f"д. {house}" if house else ""] if p]
        return ", ".join(parts) if parts else (d.get("display_name") or "")
    except Exception:
        return ""


@router.get("/inn-lookup")
async def storefront_inn_lookup(inn: str = ""):
    """Название организации по ИНН: DaData (если есть токен), иначе открытый ЕГРЮЛ."""
    inn = re.sub(r"\D", "", inn or "")
    if len(inn) not in (10, 12):
        return {"status": "error", "detail": "ИНН должен содержать 10 или 12 цифр"}
    hit = _INN_CACHE.get(inn)
    if hit and time.time() - hit["ts"] < 86400:
        return {"status": "ok", "inn": inn, "name": hit["name"],
                "kpp": hit.get("kpp", ""), "address": hit.get("address", ""),
                "ogrn": hit.get("ogrn", ""), "type": hit.get("type", ""), "cached": True}

    name, kpp, address, ogrn, ptype, vyp_tok = "", "", "", "", "", ""
    # Полные реквизиты: DaData findById/party → имя, КПП, ОГРН, юр.адрес, тип
    res = await dadata_api.validate_inn(inn)
    if res.get("status") == "found":
        name = res.get("name", "")
        kpp = res.get("kpp", "")
        address = res.get("address", "")
        ogrn = res.get("ogrn", "")
        ptype = res.get("type", "")
    if not name:
        try:
            async with httpx.AsyncClient(timeout=10, follow_redirects=True,
                                         headers={"User-Agent": "Mozilla/5.0"}) as c:
                r = await c.get(f"https://egrul.org/{inn}.json")
            if r.status_code == 200:
                d = r.json() or {}
                ul = d.get("СвЮЛ") or {}
                if ul:
                    attrs = ul.get("@attributes") or {}
                    short = ((ul.get("СвНаимЮЛСокр") or {}).get("@attributes") or {}).get("НаимСокр") or ""
                    full = ((ul.get("СвНаимЮЛ") or {}).get("@attributes") or {}).get("НаимЮЛПолн") or ""
                    name = (short or full).strip()
                    kpp = attrs.get("КПП") or ""
                    ogrn = attrs.get("ОГРН") or ""
                    ptype = "LEGAL"
                else:
                    ip = d.get("СвИП") or {}
                    attrs = ip.get("@attributes") or {}
                    fl = ((ip.get("СвФЛ") or {}).get("@attributes") or {})
                    fio = " ".join(x for x in [fl.get("Фамилия"), fl.get("Имя"), fl.get("Отчество")] if x)
                    name = f"ИП {fio}".strip() if fio else ""
                    ogrn = attrs.get("ОГРНИП") or attrs.get("ОГРН") or ""
                    ptype = "INDIVIDUAL"
        except Exception:
            name = ""
    if not name:
        try:
            # Официальный ЕГРЮЛ ФНС: POST query → токен → search-result (бесплатно, без ключа)
            async with httpx.AsyncClient(timeout=12, follow_redirects=True,
                                         headers={"User-Agent": "Mozilla/5.0"}) as c:
                r = await c.post("https://egrul.nalog.ru", data={"query": inn})
                tok = (r.json() or {}).get("t")
                if tok:
                    for _ in range(5):
                        await asyncio.sleep(0.9)
                        r2 = await c.get(f"https://egrul.nalog.ru/search-result/{tok}")
                        rows = (r2.json() or {}).get("rows") or []
                        if rows:
                            row = rows[0]
                            name = (row.get("n") or "").strip()
                            ogrn = row.get("o") or ogrn
                            # поле "p" — КПП (только у ЮЛ); "k" — тип ul/fl, "t" — токен выписки
                            p_v = str(row.get("p") or "")
                            if p_v.isdigit() and len(p_v) == 9:
                                kpp = p_v
                            address = (row.get("a") or "").strip() or address
                            ptype = "INDIVIDUAL" if len(inn) == 12 else "LEGAL"
                            vyp_tok = str(row.get("t") or "")
                            break
        except Exception:
            name = ""

    # Юр.адрес ЮЛ — из электронной выписки ФНС (PDF), ~5-15 сек при первом запросе
    if name and ptype == "LEGAL" and not address and vyp_tok:
        address = await _egrul_legal_address(vyp_tok)

    if not name:
        return {"status": "error", "detail": "Организация по этому ИНН не найдена"}
    _INN_CACHE[inn] = {"ts": time.time(), "name": name, "kpp": kpp, "address": address,
                       "ogrn": ogrn, "type": ptype}
    return {"status": "ok", "inn": inn, "name": name, "kpp": kpp,
            "address": address, "ogrn": ogrn, "type": ptype}


async def _egrul_legal_address(vyp_tok: str) -> str:
    """Юр.адрес ЮЛ из электронной выписки ЕГРЮЛ: vyp-request → vyp-status → PDF → fitz."""
    try:
        async with httpx.AsyncClient(timeout=15, follow_redirects=True,
                                     headers={"User-Agent": "Mozilla/5.0"}) as c:
            r = await c.get(f"https://egrul.nalog.ru/vyp-request/{vyp_tok}")
            vt = (r.json() or {}).get("t")
            if not vt:
                return ""
            for _ in range(15):
                await asyncio.sleep(1.0)
                st = (await c.get(f"https://egrul.nalog.ru/vyp-status/{vt}")).json() or {}
                if st.get("status") == "ready":
                    break
            else:
                return ""
            pdf = (await c.get(f"https://egrul.nalog.ru/vyp-download/{vt}")).content
        import fitz
        doc = fitz.open(stream=pdf, filetype="pdf")
        lines = [l.strip() for i in range(min(doc.page_count, 8))
                 for l in doc[i].get_text().split("\n") if l.strip()]
        for i, l in enumerate(lines):
            if l == "Адрес юридического лица":
                parts = []
                for j in range(i + 1, min(i + 8, len(lines))):
                    if lines[j].isdigit():          # номер строки таблицы — конец значения
                        break
                    parts.append(lines[j])
                return " ".join(" ".join(parts).split()).strip(" ,")
    except Exception:
        pass
    return ""


async def _notify_manager(text: str):
    """Лёгкое уведомление менеджеру в Telegram — если бот настроен."""
    if not (settings.telegram_bot_token and settings.manager_telegram_chat_id):
        return
    try:
        async with httpx.AsyncClient(timeout=6) as c:
            await c.post(
                f"https://api.telegram.org/bot{settings.telegram_bot_token}/sendMessage",
                json={"chat_id": settings.manager_telegram_chat_id, "text": text[:900]},
            )
    except Exception:
        pass


def _invoice_pdf(req, items) -> bytes:
    """Счёт на оплату — 1C-макет, см. services/invoice_generator.py."""
    return render_invoice_pdf(req, items)


class InvoiceReq(BaseModel):
    inn: str = ""
    kpp: str = ""
    name: str = ""
    address: str = ""
    phone: str = ""
    token: str = ""  # опционально — токен кабинета, чтобы счёт сохранился в «Документы»
    items: list[CartItem] = []


@router.post("/invoice")
async def storefront_invoice(req: InvoiceReq):
    """Мгновенный PDF-счёт с НДС по позициям корзины витрины."""
    inn = re.sub(r"\D", "", req.inn or "")
    if len(inn) not in (10, 12):
        return {"status": "error", "detail": "Укажите корректный ИНН (10 или 12 цифр)"}
    items = [i for i in (req.items or []) if i.name and i.price > 0 and i.qty > 0]
    if not items:
        return {"status": "error", "detail": "Корзина пуста"}
    req.inn = inn
    pdf = _invoice_pdf(req, items)
    total = sum(i.qty * i.price for i in items)
    if len(req.token) == 64 and user_store.get_profile(req.token):
        title = f"Счёт {req.name or inn} · {total:,.0f} ₽".replace(",", " ")
        user_store.add_doc(req.token, title, {"inn": inn, "name": req.name, "kpp": req.kpp,
                                              "address": req.address,
                                              "phone": req.phone, "items": [i.model_dump() for i in items]})
    try:
        user_store.add_order({
            "source": "moscow.stroiapp.ru",
            "name": req.name or f"ИНН {inn}",
            "phone": req.phone, "address": req.address, "inn": inn,
            "payment": "Безналичный счёт ООО «МИКС СТРОЙ»",
            "total": total,
            "items": [i.model_dump() for i in items],
        })
    except Exception as exc:
        print("[storefront] order log error:", exc)
    await _notify_manager(
        f"📄 Счёт: {req.name or inn} (ИНН {inn}, {req.phone or 'без телефона'}), "
        f"{len(items)} поз., {total:.0f} ₽"
    )
    return StreamingResponse(
        io.BytesIO(pdf), media_type="application/pdf",
        headers={"Content-Disposition": 'attachment; filename="invoice_stroiapp.pdf"'},
    )


@router.post("/invoice-request")
async def storefront_invoice_request(req: InvoiceReq):
    """Заявка на счёт при пустой корзине — менеджер выставит вручную."""
    inn = re.sub(r"\D", "", req.inn or "")
    if len(inn) not in (10, 12):
        return {"status": "error", "detail": "Укажите корректный ИНН (10 или 12 цифр)"}
    print(f"[storefront] invoice-request: inn={inn} name={req.name} phone={req.phone}")
    await _notify_manager(
        f"📄 Заявка на счёт: {req.name or inn} (ИНН {inn}, {req.phone or 'без телефона'}) — корзина пуста"
    )
    return {"status": "ok"}


class OrderReq(BaseModel):
    name: str = ""
    phone: str = ""
    address: str = ""
    comment: str = ""
    payment: str = "Наличные"          # Наличные | Безналичный счёт ООО «МИКС СТРОЙ»
    source: str = "moscow.stroiapp.ru"
    items: list[CartItem] = []


@router.post("/order")
async def storefront_order(req: OrderReq):
    """Оформление заказа корзины витрины (не счёт) — заявка менеджеру.
    Пишется в app.db.storefront_orders → раздел «Заказы» админки."""
    items = [i for i in (req.items or []) if i.name and i.price > 0 and i.qty > 0]
    if not items:
        raise HTTPException(status_code=400, detail="Корзина пуста")
    if not _normalize_phone(req.phone):
        raise HTTPException(status_code=422, detail="Укажите телефон для связи")
    total = sum(i.qty * i.price for i in items)
    oid = user_store.add_order({
        "source": req.source or "moscow.stroiapp.ru",
        "name": req.name or "Гость",
        "phone": req.phone,
        "address": req.address,
        "payment": req.payment or "Наличные",
        "total": total,
        "items": [i.model_dump() for i in items],
    })
    print(f"[storefront] order #{oid}: {req.name} {req.phone} total={total:.0f} pay={req.payment}")
    await _notify_manager(
        f"🛒 Заказ #{oid} ({req.source or 'moscow.stroiapp.ru'}): {req.name or '—'}, "
        f"{req.phone}, {len(items)} поз., {total:.0f} ₽ — {req.payment or 'Наличные'}"
    )
    return {"status": "ok", "order_id": oid, "total": round(total, 2)}


# ---------- Вход по SMS-коду (покупатели витрины) ----------

_SMS_CODES: dict[str, dict] = {}  # phone -> {"code": str, "ts": float, "tries": int}
_SMS_TTL = 600          # код живёт 10 минут
_SMS_RESEND = 45        # повторная отправка не чаще, чем раз в 45 сек


def _normalize_phone(phone: str) -> str:
    d = re.sub(r"\D", "", phone or "")
    if d.startswith("8") and len(d) == 11:
        d = "7" + d[1:]
    if len(d) == 10:
        d = "7" + d
    return d


class SmsSendReq(BaseModel):
    phone: str = ""


class SmsVerifyReq(BaseModel):
    phone: str = ""
    code: str = ""


@router.post("/auth/sms-code")
async def storefront_sms_code(req: SmsSendReq):
    """Отправка 4-значного кода. Без SMS-провайдера — dev-режим (код в ответе)."""
    phone = _normalize_phone(req.phone)
    if len(phone) != 11 or not phone.startswith("7"):
        return {"status": "error", "detail": "Укажите корректный номер телефона"}
    prev = _SMS_CODES.get(phone)
    if prev and time.time() - prev["ts"] < _SMS_RESEND:
        wait = int(_SMS_RESEND - (time.time() - prev["ts"]))
        return {"status": "error", "detail": f"Код уже отправлен, повторите через {wait} сек"}

    code = f"{secrets.randbelow(10000):04d}"
    _SMS_CODES[phone] = {"code": code, "ts": time.time(), "tries": 0}

    if settings.sms_api_id:
        # SMS Aero: https://smsaero.ru/api/
        try:
            async with httpx.AsyncClient(timeout=8) as c:
                r = await c.get(
                    "https://gate.smsaero.ru/v2/sms/send",
                    params={"number": phone, "text": f"Ваш код StroiApp: {code}",
                            "sign": "SMS Aero", "channel": "DIRECT"},
                    auth=(settings.sms_api_id.split("@")[0] + "@aero"
                          if "@" not in settings.sms_api_id
                          else settings.sms_api_id, ""),
                )
            if r.status_code != 200:
                return {"status": "error", "detail": "Не удалось отправить SMS"}
        except Exception:
            return {"status": "error", "detail": "SMS-сервис недоступен"}
        return {"status": "ok", "ttl": _SMS_TTL}

    print(f"[storefront] SMS code for {phone}: {code}")  # dev-режим — код в логе
    return {"status": "ok", "ttl": _SMS_TTL, "dev_code": code}


@router.post("/auth/verify")
async def storefront_auth_verify(req: SmsVerifyReq):
    """Проверка кода → токен сессии покупателя витрины."""
    phone = _normalize_phone(req.phone)
    code = re.sub(r"\D", "", req.code or "")
    rec = _SMS_CODES.get(phone)
    if not rec or time.time() - rec["ts"] > _SMS_TTL:
        return {"status": "error", "detail": "Код истёк — запросите новый"}
    rec["tries"] += 1
    if rec["tries"] > 5:
        _SMS_CODES.pop(phone, None)
        return {"status": "error", "detail": "Превышено число попыток — запросите новый код"}
    if code != rec["code"]:
        return {"status": "error", "detail": "Неверный код"}
    _SMS_CODES.pop(phone, None)
    raw = f"user:{phone}:{settings.auth_secret}"
    token = hashlib.sha256(raw.encode()).hexdigest()
    user_store.ensure_user(token, phone)
    print(f"[storefront] user logged in: {phone}")
    return {"status": "ok", "token": token, "user": {"phone": phone}}


# ---------- Списки покупателя витрины (избранное/сравнение) ----------

class UserSyncReq(BaseModel):
    token: str = ""
    wishlist: list = []
    compare: list = []


@router.get("/user/lists")
async def storefront_user_lists(token: str = ""):
    """Текущие списки пользователя с сервера (для подтягивания на новом устройстве)."""
    if len(token) != 64:
        return {"status": "error", "detail": "Неверный токен"}
    lists = user_store.get_lists(token)
    return {"status": "ok", **lists}


@router.post("/user/sync")
async def storefront_user_sync(req: UserSyncReq):
    """Слияние localStorage ↔ сервер при входе: объединяем по id товара.
    Возвращает итоговые списки, которые фронт сохраняет обратно в localStorage."""
    if len(req.token) != 64:
        return {"status": "error", "detail": "Неверный токен"}
    merged = user_store.sync_lists(req.token, req.wishlist, req.compare)
    return {"status": "ok", **merged}


# ---------- Личный кабинет (Authorization: Bearer <token>) ----------

def _bearer_token(request: Request) -> str:
    """Токен покупателя из заголовка Authorization: Bearer ... или ?token=.
    Токен валиден только если пользователь прошёл SMS-verify (есть в сторе)."""
    h = request.headers.get("authorization", "")
    tok = h[7:].strip() if h.lower().startswith("bearer ") else ""
    tok = tok or (request.query_params.get("token") or "")
    if len(tok) != 64 or not user_store.get_profile(tok):
        return ""
    return tok


class ProfileReq(BaseModel):
    name: str = ""
    email: str = ""
    inn: str = ""
    kpp: str = ""
    company: str = ""


class AddressReq(BaseModel):
    title: str = ""
    city: str = ""
    street: str = ""
    contact_name: str = ""
    contact_phone: str = ""


class SaveItemReq(BaseModel):
    title: str = ""
    payload: dict | list | str | None = None


@router.get("/user/profile")
async def storefront_profile(request: Request):
    tok = _bearer_token(request)
    if not tok:
        return {"status": "error", "detail": "Требуется авторизация"}
    return {"status": "ok", "profile": user_store.get_profile(tok)}


@router.post("/user/profile")
async def storefront_profile_save(request: Request, req: ProfileReq):
    tok = _bearer_token(request)
    if not tok:
        return {"status": "error", "detail": "Требуется авторизация"}
    return {"status": "ok", "profile": user_store.save_profile(tok, req.model_dump())}


@router.get("/user/orders")
async def storefront_orders(request: Request):
    """История заказов из OpenCart по телефону покупателя."""
    tok = _bearer_token(request)
    if not tok:
        return {"status": "error", "detail": "Требуется авторизация"}
    phone = (user_store.get_profile(tok) or {}).get("phone", "")
    orders = []
    try:
        res = await opencart_api.get_action("order/list", {"limit": 200})
        for o in (res.get("orders") or res.get("data") or []):
            op = re.sub(r"\D", "", str(o.get("telephone") or o.get("phone") or ""))
            if phone and op.endswith(phone[-10:]):
                items = o.get("products") or o.get("items") or []
                orders.append({
                    "id": o.get("order_id") or o.get("id"),
                    "date": str(o.get("date_added") or o.get("date") or "")[:10],
                    "total": float(o.get("total") or 0),
                    "status": o.get("order_status") or o.get("status") or "В обработке",
                    "items": [{"sku": it.get("model") or it.get("sku") or "",
                               "name": it.get("name") or "",
                               "qty": it.get("quantity") or it.get("qty") or 1,
                               "price": float(it.get("price") or 0)} for it in items],
                })
    except Exception as e:
        print(f"[storefront] orders fetch failed: {e}")
    orders.sort(key=lambda x: str(x.get("date") or ""), reverse=True)
    return {"status": "ok", "orders": orders}


@router.get("/user/docs")
async def storefront_docs(request: Request):
    """Список сохранённых b2b-счетов пользователя."""
    tok = _bearer_token(request)
    if not tok:
        return {"status": "error", "detail": "Требуется авторизация"}
    return {"status": "ok", "docs": user_store.list_docs(tok)}


@router.get("/user/docs/{doc_id}/pdf")
async def storefront_doc_pdf(request: Request, doc_id: int):
    """Скачивание счёта: регенерируем PDF из сохранённого payload."""
    tok = _bearer_token(request)
    if not tok:
        return {"status": "error", "detail": "Требуется авторизация"}
    doc = user_store.get_doc(tok, doc_id)
    if not doc:
        return {"status": "error", "detail": "Документ не найден"}
    p = doc["payload"] or {}
    items = [CartItem(**it) for it in (p.get("items") or []) if isinstance(it, dict)]
    pdf = _invoice_pdf(InvoiceReq(inn=p.get("inn", ""), name=p.get("name", ""),
                                  kpp=p.get("kpp", ""), address=p.get("address", ""),
                                  phone=p.get("phone", "")), items)
    return StreamingResponse(io.BytesIO(pdf), media_type="application/pdf", headers={
        "Content-Disposition": f'attachment; filename="invoice_stroiapp_{doc_id}.pdf"',
    })


@router.get("/user/calcs")
async def storefront_calcs(request: Request):
    """Сохранённые сметы из ИИ-калькулятора."""
    tok = _bearer_token(request)
    if not tok:
        return {"status": "error", "detail": "Требуется авторизация"}
    return {"status": "ok", "calcs": user_store.list_calcs(tok)}


@router.post("/user/calcs")
async def storefront_calc_save(request: Request, req: SaveItemReq):
    tok = _bearer_token(request)
    if not tok:
        return {"status": "error", "detail": "Требуется авторизация"}
    cid = user_store.add_calc(tok, req.title, req.payload)
    return {"status": "ok", "id": cid}


@router.get("/user/addresses")
async def storefront_addresses(request: Request):
    tok = _bearer_token(request)
    if not tok:
        return {"status": "error", "detail": "Требуется авторизация"}
    return {"status": "ok", "addresses": user_store.list_addresses(tok)}


@router.post("/user/addresses")
async def storefront_address_add(request: Request, req: AddressReq):
    tok = _bearer_token(request)
    if not tok:
        return {"status": "error", "detail": "Требуется авторизация"}
    if not (req.street or req.city):
        return {"status": "error", "detail": "Укажите адрес объекта"}
    return {"status": "ok", **user_store.add_address(tok, req.model_dump())}


@router.delete("/user/addresses/{address_id}")
async def storefront_address_del(request: Request, address_id: int):
    tok = _bearer_token(request)
    if not tok:
        return {"status": "error", "detail": "Требуется авторизация"}
    ok = user_store.delete_address(tok, address_id)
    return {"status": "ok" if ok else "error", "detail": "" if ok else "Адрес не найден"}


@router.get("/health")
async def storefront_health():
    """Публичный health-check: жив ли сервис и читается ли SQLite user-store."""
    db_ok = True
    try:
        user_store.init_db()
    except Exception:
        db_ok = False
    return {
        "status": "ok" if db_ok else "degraded",
        "time": datetime.utcnow().isoformat(timespec="seconds") + "Z",
        "user_store": "ok" if db_ok else "error",
    }


# ---------- v1-алиасы ----------
@v1_router.post("/upload")
async def v1_parser_upload(file: UploadFile):
    """Алиас: POST /api/v1/parser/upload — тот же пайплайн."""
    return await storefront_parser_upload(file)


@v1_router.post("/feedback")
async def v1_parser_feedback(payload: _Feedback):
    """Алиас: POST /api/v1/parser/feedback — обучение ядра."""
    return await parser_feedback(payload)
